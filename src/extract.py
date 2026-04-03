from __future__ import annotations

from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path
import re
import unicodedata
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook

from src.config import (
    BASE_DIR,
    CONFLICTOS_FUENTES_PATH,
    DIM_COMUNA_BASE_PATH,
    METADATA_PATH,
    OUTPUTS_DIR,
    PERFILADO_FUENTES_PATH,
    ensure_directories,
)

SPREADSHEETML_NS = {
    "ss": "urn:schemas-microsoft-com:office:spreadsheet",
}
SS_NAMESPACE = "{urn:schemas-microsoft-com:office:spreadsheet}"
SPECIAL_NULL_TOKENS = {
    "",
    "No Aplica",
    "No Recepcionado",
}
TEXT_SAMPLE_SIZE = 5


@dataclass(frozen=True)
class SourceContract:
    id_fuente: str
    nombre_fuente: str
    institucion: str
    url: str
    fecha_descarga: str
    formato: str
    anio_referencia: str
    variable_principal: str
    archivo_origen: str
    archivo_logico: str
    hoja: str
    skiprows: int
    observaciones: str

    @property
    def path(self) -> Path:
        return BASE_DIR / self.archivo_origen

    @property
    def uses_spreadsheetml(self) -> bool:
        return "spreadsheetml" in self.formato.lower() or self.path.suffix.lower() == ".xls"


@dataclass
class SourceDataset:
    contract: SourceContract
    raw_df: pd.DataFrame
    dataframe: pd.DataFrame
    diagnostics: dict[str, object]


def _normalize_display_text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return " ".join(str(value).split())


def _slugify(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", _normalize_display_text(value))
    ascii_only = normalized.encode("ascii", "ignore").decode("ascii")
    ascii_only = re.sub(r"[^A-Za-z0-9]+", "_", ascii_only).strip("_")
    return ascii_only.lower() or "columna"


def _normalize_name_key(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", _normalize_display_text(value))
    ascii_only = normalized.encode("ascii", "ignore").decode("ascii")
    return " ".join(ascii_only.upper().split())


def _truncate_text(value: object, max_len: int = 100) -> str:
    text = _normalize_display_text(value)
    return text if len(text) <= max_len else f"{text[: max_len - 3]}..."


def _make_unique(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    unique: list[str] = []
    for column in columns:
        count = seen.get(column, 0)
        if count == 0:
            unique.append(column)
        else:
            unique.append(f"{column}_{count + 1}")
        seen[column] = count + 1
    return unique


def _extract_indicator_code(value: object) -> str:
    text = _normalize_display_text(value)
    match = re.search(r"\b([A-Z]{2,}\d*)\b", text)
    return match.group(1).lower() if match else ""


def _excel_cell_value(value: object) -> object:
    if value is None or (isinstance(value, str) and value == "") or pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _dataframe_rows_for_excel(dataframe: pd.DataFrame) -> list[tuple[object, ...]]:
    rows: list[tuple[object, ...]] = [tuple(map(str, dataframe.columns))]
    rows.extend(
        tuple(_excel_cell_value(value) for value in row)
        for row in dataframe.itertuples(index=False, name=None)
    )
    return rows


def _workbook_matches_dataframes(
    path: Path,
    worksheets: dict[str, pd.DataFrame],
) -> bool:
    if not path.exists():
        return False

    workbook = load_workbook(path, data_only=True)
    try:
        if tuple(workbook.sheetnames) != tuple(worksheets.keys()):
            return False

        for sheet_name, dataframe in worksheets.items():
            worksheet = workbook[sheet_name]
            if worksheet.max_column != len(dataframe.columns):
                return False

            existing_rows = list(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=worksheet.max_row,
                    max_col=worksheet.max_column,
                    values_only=True,
                )
            )
            existing_rows = [
                tuple(_excel_cell_value(value) for value in row)
                for row in existing_rows
            ]
            while existing_rows and all(cell is None for cell in existing_rows[-1]):
                existing_rows.pop()

            expected_rows = _dataframe_rows_for_excel(dataframe)
            if len(existing_rows) != len(expected_rows):
                return False
            if any(tuple(existing_row) != expected_row for existing_row, expected_row in zip(existing_rows, expected_rows)):
                return False

        return True
    finally:
        workbook.close()


def _looks_like_year(value: object) -> bool:
    return bool(re.fullmatch(r"\d{4}", _normalize_display_text(value)))


def _load_metadata_contracts(metadata_path: Path = METADATA_PATH) -> dict[str, SourceContract]:
    metadata_df = pd.read_csv(metadata_path, dtype=str).fillna("")
    metadata_df["skiprows"] = metadata_df["skiprows"].astype(int)
    contracts = {
        row["id_fuente"]: SourceContract(
            id_fuente=row["id_fuente"],
            nombre_fuente=row["nombre_fuente"],
            institucion=row["institucion"],
            url=row["url"],
            fecha_descarga=row["fecha_descarga"],
            formato=row["formato"],
            anio_referencia=row["anio_referencia"],
            variable_principal=row["variable_principal"],
            archivo_origen=row["archivo_origen"],
            archivo_logico=row["archivo_logico"],
            hoja=row["hoja"],
            skiprows=int(row["skiprows"]),
            observaciones=row["observaciones"],
        )
        for _, row in metadata_df.iterrows()
    }
    return contracts


def _extract_spreadsheetml_matrix(path: Path, sheet_name: str) -> list[list[object]]:
    root = ET.fromstring(path.read_text(encoding="utf-8", errors="ignore").lstrip())
    worksheet = None
    for candidate in root.findall("ss:Worksheet", SPREADSHEETML_NS):
        if candidate.attrib.get(f"{SS_NAMESPACE}Name") == sheet_name:
            worksheet = candidate
            break
    if worksheet is None:
        raise ValueError(f"No existe la hoja {sheet_name!r} en {path.name}.")

    table = worksheet.find("ss:Table", SPREADSHEETML_NS)
    if table is None:
        raise ValueError(f"No se encontro una tabla util en {path.name}.")

    matrix: list[list[object]] = []
    for row in table.findall("ss:Row", SPREADSHEETML_NS):
        values: list[object] = []
        for cell in row.findall("ss:Cell", SPREADSHEETML_NS):
            index = cell.attrib.get(f"{SS_NAMESPACE}Index")
            if index is not None:
                while len(values) < int(index) - 1:
                    values.append(None)

            data = cell.find("ss:Data", SPREADSHEETML_NS)
            value = None if data is None else data.text
            values.append(value)

            merge_across = int(cell.attrib.get(f"{SS_NAMESPACE}MergeAcross", "0"))
            for _ in range(merge_across):
                values.append(None)
        matrix.append(values)

    width = max(len(row) for row in matrix)
    return [row + [None] * (width - len(row)) for row in matrix]


def _build_spreadsheetml_columns(
    header_row: list[object],
    descriptor_row: list[object],
) -> list[str]:
    columns: list[str] = []
    for index, (header_value, descriptor_value) in enumerate(
        zip_longest(header_row, descriptor_row, fillvalue=None),
        start=1,
    ):
        primary = _normalize_display_text(header_value)
        secondary = _normalize_display_text(descriptor_value)
        primary_upper = primary.upper()

        if primary_upper in {"CODIGO", "CODIGO COMUNA"}:
            column = "codigo_comuna"
        elif primary_upper in {"MUNICIPIO", "COMUNA", "NOMBRE COMUNA"}:
            column = "nombre_comuna"
        else:
            indicator_code = _extract_indicator_code(secondary)
            if indicator_code and _looks_like_year(primary):
                column = f"{indicator_code}_{primary}"
            elif indicator_code:
                column = indicator_code
            elif primary:
                column = _slugify(primary)
            elif secondary:
                column = _slugify(secondary)
            else:
                column = f"columna_{index}"
        columns.append(column)

    return _make_unique(columns)


def _apply_minimal_typing(df: pd.DataFrame) -> pd.DataFrame:
    typed = df.copy()
    for column in typed.columns:
        if column == "nombre_comuna":
            typed[column] = typed[column].map(_normalize_display_text).replace("", pd.NA)
            continue

        as_text = typed[column].map(_normalize_display_text)
        cleaned = as_text.replace(SPECIAL_NULL_TOKENS, pd.NA)
        numeric_candidate = pd.to_numeric(cleaned, errors="coerce")
        valid_values = cleaned.notna().sum()

        if column == "codigo_comuna":
            typed[column] = numeric_candidate.astype("Int64")
        elif valid_values > 0 and numeric_candidate.notna().sum() == valid_values:
            typed[column] = numeric_candidate
        else:
            typed[column] = cleaned
    return typed


def _read_spreadsheetml_source(contract: SourceContract) -> SourceDataset:
    matrix = _extract_spreadsheetml_matrix(contract.path, contract.hoja)
    descriptor_index = max(contract.skiprows - 1, 0)
    header_index = contract.skiprows

    header_row = matrix[header_index]
    descriptor_row = matrix[descriptor_index]
    columns = _build_spreadsheetml_columns(header_row, descriptor_row)

    data_rows = matrix[header_index + 1 :]
    raw_df = pd.DataFrame(data_rows, columns=columns).dropna(how="all").reset_index(drop=True)
    typed_df = _apply_minimal_typing(raw_df)

    special_values: dict[str, list[str]] = {}
    for column in raw_df.columns:
        values = raw_df[column].dropna().map(_normalize_display_text)
        detected = sorted({value for value in values if value in SPECIAL_NULL_TOKENS})
        if detected:
            special_values[column] = detected

    diagnostics = {
        "source_parser": "spreadsheetml_xml_2003",
        "header_descriptor_row_index": descriptor_index,
        "header_row_index": header_index,
        "special_values": special_values,
    }
    return SourceDataset(contract=contract, raw_df=raw_df, dataframe=typed_df, diagnostics=diagnostics)


def _read_excel_source(contract: SourceContract) -> SourceDataset:
    raw_df = pd.read_excel(
        contract.path,
        sheet_name=contract.hoja,
        skiprows=contract.skiprows,
    )
    raw_df.columns = [_normalize_display_text(column) for column in raw_df.columns]
    diagnostics = {
        "source_parser": "pandas_read_excel",
    }
    return SourceDataset(contract=contract, raw_df=raw_df.copy(), dataframe=raw_df.copy(), diagnostics=diagnostics)


def read_source(source_id: str, contracts: dict[str, SourceContract] | None = None) -> SourceDataset:
    contracts = contracts or _load_metadata_contracts()
    contract = contracts[source_id]
    if contract.uses_spreadsheetml:
        return _read_spreadsheetml_source(contract)
    return _read_excel_source(contract)


def read_all_sources() -> dict[str, SourceDataset]:
    contracts = _load_metadata_contracts()
    return {
        source_id: read_source(source_id, contracts=contracts)
        for source_id in contracts
    }


def _text_columns(df: pd.DataFrame) -> list[str]:
    return [
        column
        for column in df.columns
        if pd.api.types.is_object_dtype(df[column]) or pd.api.types.is_string_dtype(df[column])
    ]


def _sample_unique_text_values(series: pd.Series, sample_size: int = TEXT_SAMPLE_SIZE) -> str:
    cleaned = [
        _truncate_text(value)
        for value in series.dropna().map(_normalize_display_text)
        if _normalize_display_text(value)
    ]
    unique_values = list(dict.fromkeys(cleaned))
    return " | ".join(unique_values[:sample_size])


def _profile_source(source_id: str, dataset: SourceDataset) -> tuple[dict[str, object], list[dict[str, object]]]:
    df = dataset.dataframe
    summary_row = {
        "source_id": source_id,
        "nombre_fuente": dataset.contract.nombre_fuente,
        "archivo_origen": dataset.contract.archivo_origen,
        "archivo_logico": dataset.contract.archivo_logico,
        "formato": dataset.contract.formato,
        "hoja": dataset.contract.hoja,
        "skiprows": dataset.contract.skiprows,
        "filas": len(df),
        "columnas": len(df.columns),
        "duplicados_fila": int(df.duplicated().sum()),
        "filas_totalmente_nulas": int(df.isna().all(axis=1).sum()),
        "nombres_columnas": " | ".join(map(str, df.columns)),
        "parser": dataset.diagnostics["source_parser"],
    }

    column_rows: list[dict[str, object]] = []
    for column in df.columns:
        text_sample = ""
        if column in _text_columns(df):
            text_sample = _sample_unique_text_values(df[column])

        column_rows.append(
            {
                "source_id": source_id,
                "nombre_fuente": dataset.contract.nombre_fuente,
                "archivo_origen": dataset.contract.archivo_origen,
                "columna": column,
                "dtype": str(df[column].dtype),
                "nulos": int(df[column].isna().sum()),
                "porcentaje_nulos": round(float(df[column].isna().mean()) * 100, 2),
                "valores_unicos_no_nulos": int(df[column].nunique(dropna=True)),
                "muestra_valores_unicos_texto": text_sample,
            }
        )

    return summary_row, column_rows


def _load_dim_comuna_base() -> pd.DataFrame:
    dim_base = pd.read_csv(DIM_COMUNA_BASE_PATH, dtype={"codigo_comuna": "Int64"})
    dim_base["nombre_normalizado"] = dim_base["nombre_comuna"].map(_normalize_name_key)
    return dim_base


def _code_column_for_source(dataset: SourceDataset) -> str | None:
    if "codigo_comuna" in dataset.dataframe.columns:
        return "codigo_comuna"
    for candidate in ("Codigo", "Código", "Código comuna", "CODIGO", "Código región"):
        if candidate in dataset.dataframe.columns:
            return candidate
    return None


def _name_column_for_source(dataset: SourceDataset) -> str | None:
    if "nombre_comuna" in dataset.dataframe.columns:
        return "nombre_comuna"
    for candidate in ("Nombre comuna", "Comuna", "MUNICIPIO"):
        if candidate in dataset.dataframe.columns:
            return candidate
    return None


def _codes_in_scope(dataset: SourceDataset, dim_base: pd.DataFrame) -> tuple[int, int]:
    code_column = _code_column_for_source(dataset)
    if code_column is None:
        return 0, 0
    codes = pd.to_numeric(dataset.raw_df[code_column], errors="coerce").dropna().astype(int)
    in_scope = int(codes.isin(dim_base["codigo_comuna"].dropna().astype(int)).sum())
    return len(codes), in_scope


def _exact_name_differences(dataset: SourceDataset, dim_base: pd.DataFrame) -> pd.DataFrame:
    code_column = _code_column_for_source(dataset)
    name_column = _name_column_for_source(dataset)
    if code_column is None or name_column is None:
        return pd.DataFrame(columns=["codigo_comuna", "nombre_comuna", "nombre_fuente"])

    comparison = dataset.raw_df[[code_column, name_column]].copy()
    comparison["codigo_comuna"] = pd.to_numeric(comparison[code_column], errors="coerce").astype("Int64")
    comparison["nombre_fuente"] = comparison[name_column].map(_normalize_display_text)
    comparison = comparison.dropna(subset=["codigo_comuna"])
    comparison = comparison[comparison["codigo_comuna"].isin(dim_base["codigo_comuna"])]
    merged = dim_base.merge(comparison[["codigo_comuna", "nombre_fuente"]], on="codigo_comuna", how="left")
    return merged[merged["nombre_comuna"] != merged["nombre_fuente"]][["codigo_comuna", "nombre_comuna", "nombre_fuente"]]


def _note_like_rows(dataset: SourceDataset) -> pd.DataFrame:
    df = dataset.raw_df
    return df[df.isna().sum(axis=1) >= len(df.columns) - 1]


def _special_value_counts(dataset: SourceDataset) -> dict[str, int]:
    counts: dict[str, int] = {}
    for column in dataset.raw_df.columns:
        values = dataset.raw_df[column].dropna().map(_normalize_display_text)
        for token in SPECIAL_NULL_TOKENS - {""}:
            token_count = int((values == token).sum())
            if token_count:
                counts[token] = counts.get(token, 0) + token_count
    return counts


def _format_examples(rows: pd.DataFrame, left_col: str, right_col: str, limit: int = 5) -> str:
    examples = []
    for _, row in rows.head(limit).iterrows():
        examples.append(f"{row['codigo_comuna']}: {row[left_col]} -> {row[right_col]}")
    return "; ".join(examples)


def _build_conflicts_markdown(datasets: dict[str, SourceDataset]) -> tuple[str, list[str]]:
    dim_base = _load_dim_comuna_base()
    highlights: list[str] = []
    lines = [
        "# Conflictos detectados en Fase 2",
        "",
        "## Resumen ejecutivo",
    ]

    lines.append("- Las cuatro fuentes se leen desde `data/raw/` usando `metadata_fuentes.csv` como contrato operativo.")
    lines.append("- Las fuentes A y B requirieron parser especial para SpreadsheetML/XML 2003; C y D se leen con `pandas.read_excel()` respetando `hoja` y `skiprows` del metadata.")
    lines.append("- El diagnostico confirma que la llave estable para la unidad de analisis debe ser `codigo_comuna`; los nombres de comuna cambian en mayusculas, tildes y estilo de escritura entre fuentes.")
    lines.append("- Persisten filas fuera de alcance, filas de notas/totales y celdas especiales que deben resolverse en Fase 3/Fase 5.")
    lines.append("")
    lines.append("## Conflictos transversales")

    transversales: list[str] = []
    for source_id, dataset in datasets.items():
        name_differences = _exact_name_differences(dataset, dim_base)
        if not name_differences.empty:
            example_text = _format_examples(name_differences, "nombre_comuna", "nombre_fuente")
            transversales.append(
                f"- Fuente {source_id}: hay {len(name_differences)} diferencias de nombre exacto frente a `dim_comuna_base.csv`. Ejemplos: {example_text}."
            )

    transversales.append(
        "- Priorizar `codigo_comuna` sobre `nombre_comuna` en joins y validaciones; las variantes de nombre son compatibles solo despues de normalizacion de mayusculas y tildes."
    )
    transversales.append(
        "- Todas las fuentes exceden el alcance final de la Provincia de Santiago: deben filtrarse usando `dim_comuna_base.csv` antes de integrar."
    )
    lines.extend(transversales)
    lines.append("")

    for source_id, dataset in datasets.items():
        total_codes, in_scope = _codes_in_scope(dataset, dim_base)
        out_of_scope = total_codes - in_scope
        note_rows = _note_like_rows(dataset)
        special_counts = _special_value_counts(dataset)
        name_differences = _exact_name_differences(dataset, dim_base)
        highlights.append(
            f"Fuente {source_id}: {len(dataset.dataframe)} filas leidas, {in_scope} comunas objetivo cubiertas y {out_of_scope} filas/codigos fuera de alcance."
        )

        lines.append(f"## Fuente {source_id} - {dataset.contract.nombre_fuente}")
        lines.append(f"- Archivo: `{dataset.contract.archivo_origen}`")
        lines.append(f"- Parser aplicado: `{dataset.diagnostics['source_parser']}`")
        lines.append(
            f"- Lectura operativa: hoja `{dataset.contract.hoja}` con `skiprows={dataset.contract.skiprows}`."
        )
        lines.append(
            f"- Cobertura detectada por codigo: {total_codes} codigos validos; {in_scope} comunas de la Provincia de Santiago y {out_of_scope} registros fuera de alcance."
        )

        if dataset.contract.uses_spreadsheetml:
            lines.append(
                "- El archivo trae encabezado extendido previo a la tabla util. Se reconstruyeron columnas utilizables a partir de la fila descriptiva y la fila de encabezado."
            )

        if special_counts:
            special_text = ", ".join(f"{token}={count}" for token, count in sorted(special_counts.items()))
            lines.append(
                f"- Se detectaron valores especiales no numericos en columnas operativamente numericas: {special_text}."
            )

        if not note_rows.empty:
            lines.append(
                f"- Se detectaron {len(note_rows)} filas en blanco o de nota que deben excluirse antes de tipificar o integrar."
            )

        if source_id == "D":
            total_country_rows = int(
                (pd.to_numeric(dataset.raw_df.get("Código comuna"), errors="coerce").fillna(-1) == 0).sum()
            )
            if total_country_rows:
                lines.append(
                    f"- La hoja incluye {total_country_rows} fila total/agrupadora (`País`) que no pertenece a la unidad de analisis comunal."
                )

        if source_id == "C":
            lines.append(
                "- La variable de pobreza viene como proporcion entre 0 y 1, no como porcentaje 0-100. Esa conversion debe tratarse en la fase de transformacion."
            )

        if not name_differences.empty:
            lines.append(
                f"- Diferencias de nombre exacto respecto a la dimension base: {len(name_differences)}. Ejemplos: {_format_examples(name_differences, 'nombre_comuna', 'nombre_fuente')}."
            )
        else:
            lines.append(
                "- No se detectaron diferencias de nombre exacto dentro del alcance final una vez que se compara por `codigo_comuna`."
            )

        lines.append("")

    lines.append("## Listo para Fase 3 y Fase 5")
    lines.append("- Filtrar las cuatro fuentes por `codigo_comuna` usando `dim_comuna_base.csv`.")
    lines.append("- Excluir filas de notas, filas totalmente vacias y totales antes de convertir tipos.")
    lines.append("- Convertir valores especiales como `No Aplica` y `No Recepcionado` a nulos controlados.")
    lines.append("- Mantener los nombres de comuna solo como apoyo de validacion y presentacion; no como llave de integracion.")

    return "\n".join(lines) + "\n", highlights


def export_profiles(
    datasets: dict[str, SourceDataset],
    output_path: Path = PERFILADO_FUENTES_PATH,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    summary_rows: list[dict[str, object]] = []
    column_rows: list[dict[str, object]] = []

    for source_id, dataset in datasets.items():
        summary_row, per_column_rows = _profile_source(source_id, dataset)
        summary_rows.append(summary_row)
        column_rows.extend(per_column_rows)

    summary_df = pd.DataFrame(summary_rows)
    columns_df = pd.DataFrame(column_rows)
    worksheets = {
        "resumen": summary_df,
        "columnas": columns_df,
    }
    if _workbook_matches_dataframes(output_path, worksheets):
        return summary_df, columns_df

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="resumen", index=False)
        columns_df.to_excel(writer, sheet_name="columnas", index=False)
    return summary_df, columns_df


def export_conflicts(
    datasets: dict[str, SourceDataset],
    output_path: Path = CONFLICTOS_FUENTES_PATH,
) -> tuple[str, list[str]]:
    markdown, highlights = _build_conflicts_markdown(datasets)
    output_path.write_text(markdown, encoding="utf-8")
    return markdown, highlights


def run_phase_2_profile() -> dict[str, object]:
    ensure_directories()
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

    datasets = read_all_sources()
    summary_df, columns_df = export_profiles(datasets)
    _, highlights = export_conflicts(datasets)

    return {
        "datasets": datasets,
        "summary_df": summary_df,
        "columns_df": columns_df,
        "profile_output_path": PERFILADO_FUENTES_PATH,
        "conflicts_output_path": CONFLICTOS_FUENTES_PATH,
        "highlights": highlights,
    }
